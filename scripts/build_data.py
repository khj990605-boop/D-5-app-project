"""
뉴비키 데이터 만들기: data/회계법인_용어집.xlsx  →  data.js

사용법 (저장소 폴더에서)
    pip install openpyxl
    python scripts/build_data.py

엑셀을 고친 뒤 이 스크립트를 다시 실행하면 data.js가 새로 만들어져요.
index.html은 data.js를 읽어서 화면을 그려요.
"""
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "회계법인_용어집.xlsx"
OUT = ROOT / "data.js"

LEVEL = {"입사 첫 주": "week1", "한 달 안에": "month1", "알아두면 좋아요": "later"}
REVIEW = {"검수 완료": "approved", "검토 중": "pending", "수정 필요": "needs_fix"}
ANON = {"", "미기재", "익명", None}

errors = []


def norm(s):
    return re.sub(r"\s+", "", str(s)).lower()


def split_list(s):
    return [x.strip() for x in str(s or "").split(",") if x.strip()]


def to_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if v:
        return str(v)[:10]
    return date.today().strftime("%Y-%m-%d")


def split_examples(s):
    """'“문장1” / “문장2”' → ['문장1', '문장2']  (한 묶음에 따옴표가 여러 개면 ' — '로 이어 붙임)"""
    out = []
    for chunk in re.split(r"\s+/\s+", str(s or "")):
        quoted = re.findall(r"[“\"](.+?)[”\"]", chunk)
        text = " — ".join(q.strip() for q in quoted) if quoted else chunk.strip(" “”\"")
        if text:
            out.append(text)
    return out


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # ── 카테고리 시트
    cats, cat_by_name = [], {}
    cs = wb["카테고리"]
    for name, cid, emoji, color, desc in cs.iter_rows(min_row=2, max_col=5, values_only=True):
        if not name:
            continue
        c = {"id": cid, "name": name, "color": color or "gray", "desc": desc or ""}
        if emoji:
            c["emoji"] = emoji
        cats.append(c)
        cat_by_name[name] = c

    # ── 용어집 시트 (머리글 이름으로 열을 찾아서, 열 순서가 바뀌어도 동작)
    ws = wb["용어집"]
    head = [str(c.value).strip() if c.value else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    need = ["용어명", "용어의 사내 정의", "업무 상황별 사용 예시", "관련 용어 및 연관 키워드",
            "최종 수정일", "등록자", "출처 / 비고", "id", "표시 이름", "별칭", "카테고리",
            "난이도", "검수 상태", "많이 찾은 말 순위"]
    missing = [h for h in need if h not in col]
    if missing:
        sys.exit(f"엑셀 머리글을 찾을 수 없어요: {missing}")

    # 쉽게 말하면 · 선배 팁 · 검수자 열은 있으면 쓰고, 없어도 돼요
    rows = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        g = lambda h, row=row: row[col[h]] if h in col else None
        if not g("용어명") and not g("표시 이름"):
            continue
        rows.append((r, g))

    terms, details, popular, seen = [], {}, [], set()
    for r, g in rows:
        tid = str(g("id") or "").strip()
        name = str(g("표시 이름") or g("용어명")).strip()
        if not re.fullmatch(r"[a-z0-9]+", tid):
            errors.append(f"{r}행 '{name}': id는 영어 소문자와 숫자만 쓸 수 있어요 (지금: '{tid}')")
        if tid in seen:
            errors.append(f"{r}행 '{name}': id '{tid}'가 다른 행과 겹쳐요")
        seen.add(tid)
        cat = cat_by_name.get(str(g("카테고리") or "").strip())
        if not cat:
            errors.append(f"{r}행 '{name}': 카테고리 '{g('카테고리')}'가 「카테고리」 시트에 없어요")
        level = LEVEL.get(str(g("난이도") or "").strip())
        if not level:
            errors.append(f"{r}행 '{name}': 난이도는 {list(LEVEL)} 중 하나여야 해요")
        status = REVIEW.get(str(g("검수 상태") or "검토 중").strip(), "pending")
        by = str(g("등록자") or "").strip()
        updated = to_date(g("최종 수정일"))
        terms.append({
            "id": tid, "name": name, "category": cat["id"] if cat else "",
            "level": level or "later", "definition": str(g("용어의 사내 정의") or "").strip(),
            "aliases": split_list(g("별칭")),
            "createdBy": "익명 선배" if by in ANON else by,
            "createdAt": updated, "updatedAt": updated,
            "reviewStatus": status, "reviewedBy": str(g("검수자") or "").strip() or None,
            "source": str(g("출처 / 비고") or "").strip(),
            "views": 0,
            "_related_raw": split_list(g("관련 용어 및 연관 키워드")),
            "_examples": split_examples(g("업무 상황별 사용 예시")),
            "_easy": str(g("쉽게 말하면") or "").strip(),
            "_tip": str(g("선배 팁") or "").strip(),
        })
        rank = g("많이 찾은 말 순위")
        if rank not in (None, ""):
            popular.append((float(rank), tid))

    # 이름·별칭 → 용어 (관련 용어를 id로 연결할 때 사용)
    lookup = {}
    for t in terms:
        for k in [t["name"], *t["aliases"]]:
            lookup.setdefault(norm(k), []).append(t)

    for t in terms:
        own = {norm(k) for k in [t["name"], *t["aliases"]]}
        related, keywords = [], []
        for word in t.pop("_related_raw"):
            n = norm(word)
            if n in own:
                continue  # 자기 자신의 다른 이름은 건너뜀
            hits = [x for x in lookup.get(n, []) if x["id"] != t["id"]]
            if hits:
                # 같은 이름이 둘 이상이면 (예: PM) 같은 카테고리를 우선
                hits.sort(key=lambda x: x["category"] != t["category"])
                if hits[0]["id"] not in related:
                    related.append(hits[0]["id"])
            elif word not in keywords:
                keywords.append(word)
        by = t["createdBy"]
        easy, tip = t.pop("_easy"), t.pop("_tip")
        details[t["id"]] = {
            "examples": [{"where": "", "text": e, "by": by} for e in t.pop("_examples")],
            "related": related,
            "keywords": keywords,
            "history": [{"date": t["updatedAt"], "by": by, "what": "용어집에 등록했어요"}],
        }
        if easy:
            details[t["id"]]["easy"] = easy
        if tip:
            details[t["id"]]["tip"] = tip

    # 카테고리 카드에 보여줄 대표 용어 (첫 주 필수 → 한 달 → 나중 순)
    order = {"week1": 0, "month1": 1, "later": 2}
    for c in cats:
        mine = sorted([t for t in terms if t["category"] == c["id"]], key=lambda t: order[t["level"]])
        c["sample"] = [t["name"] for t in mine[:5]]
        c["count"] = len(mine)

    if errors:
        print("엑셀에서 고칠 곳이 있어요:")
        for e in errors:
            print("  -", e)
        sys.exit(1)

    popular_ids = [tid for _, tid in sorted(popular)][:3]
    if not popular_ids:
        popular_ids = [t["id"] for t in terms[:3]]

    js = (
        "/* 이 파일은 scripts/build_data.py가 자동으로 만들어요. 직접 고치지 말고 엑셀을 고친 뒤 다시 실행하세요.\n"
        f"   원본: data/{XLSX.name} · 만든 시각: {datetime.now():%Y-%m-%d %H:%M} · 용어 {len(terms)}개 */\n"
        f"const CATEGORIES = {json.dumps(cats, ensure_ascii=False, indent=2)};\n\n"
        f"const TERMS = {json.dumps(terms, ensure_ascii=False, indent=2)};\n\n"
        f"const DETAILS = {json.dumps(details, ensure_ascii=False, indent=2)};\n\n"
        f"const POPULAR = {json.dumps(popular_ids, ensure_ascii=False)};\n\n"
        "/* 요청 수는 앱이 구글 시트에서 직접 받아요 */\n"
        "const REQUESTS = [];\n"
    )
    OUT.write_text(js, encoding="utf-8")
    print(f"완료: {OUT.name} (용어 {len(terms)}개, 카테고리 {len(cats)}개, 많이 찾은 말 {popular_ids})")


if __name__ == "__main__":
    main()
